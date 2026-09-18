import streamlit as st
import os
import re
import smtplib
import time
import imaplib
import json
import hashlib
import string
import email
from email.header import decode_header
from email.utils import parseaddr
from filelock import FileLock
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.application import MIMEApplication
import openpyxl

# 獨立導入資安模組 (維持檔案分離與模組化架構)
from security import (
    verify_password, generate_salted_password_hash, is_legacy_hash,
    scraping_detector, log_security_event, get_recent_security_logs,
    check_session_timeout
)

try:
    import pandas as pd
except ImportError:
    pd = None

# ==========================================
# 網頁基本設定
# ==========================================
st.set_page_config(page_title="公關寄信系統", layout="wide")

USERS_FILE = "users.json"
USERS_LOCK = "users.json.lock"
PROJECTS_FILE = "projects.json"
PROJECTS_LOCK = "projects.json.lock"

# ==========================================
# 郵件解碼輔助功能
# ==========================================
def decode_str(s):
    """安全解析信件標題中的編碼字元"""
    if not s: return ""
    try:
        value, charset = decode_header(s)[0]
        if charset:
            return value.decode(charset)
        elif isinstance(value, bytes):
            return value.decode('utf-8', errors='ignore')
        return str(value)
    except Exception:
        return str(s)

# ==========================================
# 範本安全處理器 (Safe Formatter)
# ==========================================
class SafeFormatter(string.Formatter):
    def get_value(self, key, args, kwargs):
        if isinstance(key, str):
            return kwargs.get(key, f"{{{key}}}")
        return super().get_value(key, args, kwargs)

def safe_format_template(template, context_dict):
    if not template: return ""
    try:
        formatter = SafeFormatter()
        return formatter.format(template, **context_dict)
    except Exception:
        res = str(template)
        for k, v in context_dict.items():
            res = res.replace(f"{{{k}}}", str(v))
        return res

# ==========================================
# 資料庫與系統安全函式
# ==========================================
def safe_replace_file(src, dst, max_retries=5):
    for attempt in range(max_retries):
        try:
            os.replace(src, dst)
            return
        except PermissionError:
            if attempt == max_retries - 1: raise
            time.sleep(0.02)

def load_users():
    with FileLock(USERS_LOCK, timeout=10):
        if not os.path.exists(USERS_FILE):
            # 預設管理員帳號使用 PBKDF2 加鹽加密
            default_users = {"admin": {"password": generate_salted_password_hash("admin123"), "role": "admin", "real_name": "系統管理員"}}
            temp_file = f"{USERS_FILE}.tmp"
            with open(temp_file, "w", encoding="utf-8") as f:
                json.dump(default_users, f, ensure_ascii=False, indent=4)
            safe_replace_file(temp_file, USERS_FILE)
            return default_users
        with open(USERS_FILE, "r", encoding="utf-8") as f:
            try: return json.load(f)
            except Exception: return {}

def save_users(users_data):
    with FileLock(USERS_LOCK, timeout=10):
        temp_file = f"{USERS_FILE}.tmp"
        with open(temp_file, "w", encoding="utf-8") as f:
            json.dump(users_data, f, ensure_ascii=False, indent=4)
        safe_replace_file(temp_file, USERS_FILE)

def load_projects():
    with FileLock(PROJECTS_LOCK, timeout=10):
        if not os.path.exists(PROJECTS_FILE): return {}
        with open(PROJECTS_FILE, "r", encoding="utf-8") as f:
            try: return json.load(f)
            except Exception: return {}

def save_projects(projects_data):
    with FileLock(PROJECTS_LOCK, timeout=10):
        temp_file = f"{PROJECTS_FILE}.tmp"
        with open(temp_file, "w", encoding="utf-8") as f:
            json.dump(projects_data, f, ensure_ascii=False, indent=4)
        safe_replace_file(temp_file, PROJECTS_FILE)

def extract_email(contact_str):
    if not contact_str or str(contact_str).lower() in ("nan", "none", "null"): return None
    match = re.search(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', str(contact_str))
    return match.group(0) if match else None

def read_excel_data(uploaded_file):
    if pd is not None:
        try:
            df = pd.read_excel(uploaded_file, sheet_name="全部彙總清單")
            return df.to_dict(orient="records")
        except Exception:
            try:
                df = pd.read_excel(uploaded_file, sheet_name=0)
                return df.to_dict(orient="records")
            except Exception: pass
    try:
        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
        sheet_name = "全部彙總清單" if "全部彙總清單" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows: return []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        data = []
        for row in rows[1:]:
            if any(c is not None for c in row):
                data.append({headers[i]: (row[i] if i < len(row) and row[i] is not None else "") for i in range(len(headers))})
        return data
    except Exception as e:
        st.error(f"❌ Excel 讀取失敗：{e}")
        return []

DEFAULT_TEMPLATE = """尊敬的 {企業／贊助單位} 貴賓與團隊 您好：\n\n我們是來自台灣的高中機器人競賽團隊【{team_name}】。\n\n期盼能有機會向貴公司爭取支持與合作（{說明與贊助契機}）。隨信檢附專屬企劃書：\n▶ 專屬企劃書：{pdf_filename}\n\n敬祝 商祺\n\n【{team_name}】公關團隊\n聯絡人：{contact_person}\n聯絡電話：{contact_phone}"""

# ==========================================
# 狀態初始化與 30 分鐘會話超時檢查
# ==========================================
if "logged_in" not in st.session_state:
    st.session_state.logged_in = False
    st.session_state.username = ""
    st.session_state.role = ""
    st.session_state.real_name = ""
    st.session_state.last_active = time.time()
if "page" not in st.session_state: st.session_state.page = "home"
if "current_project" not in st.session_state: st.session_state.current_project = None
if "gmail_account" not in st.session_state: st.session_state.gmail_account = ""
if "gmail_password" not in st.session_state: st.session_state.gmail_password = ""

# 檢查會話閒置超時 (30 分鐘無操作自動銷毀 Session)
if st.session_state.logged_in:
    if check_session_timeout(st.session_state.get("last_active", 0), timeout_seconds=1800):
        log_security_event("SESSION_TIMEOUT", st.session_state.username, "EXPIRED", "閒置超過 30 分鐘自動登出")
        st.session_state.logged_in = False
        st.session_state.username = ""
        st.sidebar.error("⚠️ 會話因閒置超過 30 分鐘已自動安全登出，請重新登入！")
        st.rerun()
    st.session_state.last_active = time.time()

# ==========================================
# 登入介面 (整合 security.py PBKDF2 與日誌)
# ==========================================
if not st.session_state.logged_in:
    st.title("🔐 公關寄信系統")
    st.markdown("請輸入您的專屬帳號與密碼以登入系統。")
    with st.form("login_form"):
        login_user = st.text_input("帳號").strip()
        login_pwd = st.text_input("密碼", type="password").strip()
        if st.form_submit_button("登入", type="primary"):
            users_db = load_users()
            if login_user in users_db and verify_password(users_db[login_user]["password"], login_pwd):
                st.session_state.logged_in = True
                st.session_state.username, st.session_state.role = login_user, users_db[login_user]["role"]
                st.session_state.real_name = users_db[login_user].get("real_name", login_user)
                st.session_state.last_active = time.time()
                
                # 自動對舊版 SHA-256 密碼執行 PBKDF2 加鹽無感升級
                if is_legacy_hash(users_db[login_user]["password"]):
                    users_db[login_user]["password"] = generate_salted_password_hash(login_pwd)
                    save_users(users_db)
                    log_security_event("CREDENTIAL_UPGRADE", login_user, "SUCCESS", "密碼無感升級為 PBKDF2 加鹽雜湊")
                
                log_security_event("LOGIN_SUCCESS", login_user, "SUCCESS", f"登入成功 ({users_db[login_user]['role']})")
                st.rerun()
            else:
                log_security_event("LOGIN_FAILURE", login_user if login_user else "GUEST", "FAILURE", "登入密碼驗證失敗")
                st.error("⚠️ 帳號或密碼錯誤！")
    st.stop()

# ==========================================
# 側邊欄：動態計算未讀訊息與導覽
# ==========================================
projects_db = load_projects()

unread_count = 0
for p_data in projects_db.values():
    for reply in p_data.get("replies", []):
        if not reply.get("read", True):
            unread_count += 1

st.sidebar.markdown(f"👤 登入者：**{st.session_state.real_name}**")
if st.sidebar.button("🚪 登出", use_container_width=True):
    log_security_event("LOGOUT", st.session_state.username, "SUCCESS", "使用者手動登出")
    st.session_state.logged_in = False
    st.rerun()
st.sidebar.divider()

nav_options = ["🏠 專案與寄信區"]
nav_options.append(f"📥 收件與回信匣 {'🔴' if unread_count > 0 else ''}")
if st.session_state.role == "admin":
    nav_options.append("⚙️ 系統後台管理")

app_mode = st.sidebar.radio("📌 系統功能導覽", nav_options)
st.sidebar.divider()

# ==========================================
# 模式 A：系統後台管理 (整合資安稽核日誌)
# ==========================================
if "⚙️ 系統後台管理" in app_mode:
    st.title("⚙️ 系統後台管理")
    tab1, tab2, tab3 = st.tabs(["👥 帳號管理", "📊 團隊寄件總覽", "🛡️ 資安稽核日誌"])
    
    with tab1:
        st.subheader("建立新帳號")
        with st.form("add_user_form"):
            c1, c2, c3 = st.columns(3)
            n_usr = c1.text_input("登入帳號").strip()
            n_name = c2.text_input("成員姓名").strip()
            n_pwd = c3.text_input("預設密碼", type="password").strip()
            n_role = st.selectbox("帳號權限", ["user (一般)", "admin (管理員)"])
            if st.form_submit_button("新增帳號", type="primary"):
                users_db = load_users()
                if n_usr in users_db: st.error("帳號已存在！")
                elif not n_usr or not n_pwd: st.error("不得為空！")
                else:
                    users_db[n_usr] = {
                        "password": generate_salted_password_hash(n_pwd), 
                        "role": "admin" if "admin" in n_role else "user", 
                        "real_name": n_name or n_usr
                    }
                    save_users(users_db)
                    log_security_event("USER_CREATE", st.session_state.username, "SUCCESS", f"建立帳號 {n_usr} ({n_name})")
                    st.success(f"成功建立帳號：{n_name}")
        st.table([{"帳號": u, "姓名": d.get("real_name", u), "權限": d.get("role", "user")} for u, d in load_users().items()])

    with tab2:
        st.subheader("📂 專案進度與紀錄總覽")
        for p_name, p_data in projects_db.items():
            sent_list = p_data.get("sent_companies", [])
            with st.expander(f"📁 {p_name} (寄出 {len(sent_list)} 封)"):
                for record in sent_list:
                    if isinstance(record, dict):
                        st.markdown(f"- **{record.get('company', '?')}** (Email: {record.get('email', '未記錄')} | 寄件: {record.get('sender', '?')})")
                    else:
                        st.markdown(f"- **{record}** (早期紀錄)")

    with tab3:
        st.subheader("🛡️ 系統資安稽核日誌 (security_audit.log)")
        st.markdown("顯示全系統最近的認證事件、資料存取與外網異常擷取警報：")
        logs = get_recent_security_logs(100)
        if not logs:
            st.info("目前尚無資安稽核紀錄。")
        else:
            st.code("\n".join(logs), language="log")

# ==========================================
# 模式 B：收件與回信匣 (強制掃描與除錯版)
# ==========================================
elif "📥 收件與回信匣" in app_mode:
    st.title("📥 廠商回信與通知中心")
    st.markdown("系統會掃描團隊信箱，將**已經寄出過企劃書的廠商 Email** 來信自動拉取至此。")
    
    st.subheader("1. 郵件伺服器認證")
    col1, col2, col3 = st.columns([3, 3, 2])
    test_email = col1.text_input("團隊 Gmail 信箱", value=st.session_state.gmail_account).strip()
    test_pwd = col2.text_input("應用程式密碼", value=st.session_state.gmail_password, type="password").strip()
    reply_folder = col3.text_input("自動歸檔資料夾名稱", value="FRC_Replies", help="請盡量使用英文，以防編碼錯誤")
    
    if st.button("🔄 強制掃描近期回信", type="primary", use_container_width=True):
        if not test_email or not test_pwd:
            st.error("請輸入信箱與應用程式密碼！")
        else:
            with st.spinner("🚀 正在強制掃描 Gmail 近 3 天內所有信件，請稍候..."):
                try:
                    mail = imaplib.IMAP4_SSL("imap.gmail.com", timeout=10)
                    mail.login(test_email, test_pwd)
                    
                    status, _ = mail.select(reply_folder)
                    if status != 'OK': mail.create(reply_folder)
                    
                    mail.select("INBOX")
                    
                    date_limit = (pd.Timestamp.now() - pd.Timedelta(days=3)).strftime("%d-%b-%Y")
                    status, messages = mail.search(None, f'(SINCE "{date_limit}")')
                    
                    scanned_emails = []
                    target_emails = []

                    if status == "OK" and messages[0]:
                        msg_nums = messages[0].split()
                        
                        sent_map = {}
                        for p_name, p_data in projects_db.items():
                            if "replies" not in p_data: p_data["replies"] = []
                            for record in p_data.get("sent_companies", []):
                                if isinstance(record, dict) and record.get("email"):
                                    clean_email = str(record["email"]).strip().lower()
                                    sent_map[clean_email] = (p_name, record.get("company"))
                                    target_emails.append(clean_email)

                        new_reply_count = 0
                        if len(msg_nums) > 100: msg_nums = msg_nums[-100:]

                        for num in msg_nums:
                            res, header_data = mail.fetch(num, "(BODY.PEEK[HEADER.FIELDS (FROM SUBJECT)])")
                            for response_part in header_data:
                                if isinstance(response_part, tuple):
                                    header_msg = email.message_from_bytes(response_part[1])
                                    from_header = decode_str(header_msg.get("From"))
                                    _, addr = parseaddr(from_header)
                                    addr_lower = str(addr).strip().lower()
                                    
                                    subject_check = decode_str(header_msg.get("Subject"))
                                    scanned_emails.append(f"{addr_lower} (主旨: {subject_check})")
                                    
                                    if addr_lower in sent_map:
                                        proj_name, company_name = sent_map[addr_lower]
                                        
                                        is_duplicate = False
                                        for existing_reply in projects_db[proj_name]["replies"]:
                                            if existing_reply["subject"] == subject_check and existing_reply["email"] == addr_lower:
                                                is_duplicate = True
                                                break
                                                
                                        if not is_duplicate:
                                            res, full_msg_data = mail.fetch(num, "(RFC822)")
                                            for full_response_part in full_msg_data:
                                                if isinstance(full_response_part, tuple):
                                                    msg = email.message_from_bytes(full_response_part[1])
                                                    
                                                    body = "無法解析文字內容"
                                                    if msg.is_multipart():
                                                        for part in msg.walk():
                                                            if part.get_content_type() == "text/plain":
                                                                try: body = part.get_payload(decode=True).decode('utf-8', errors='ignore'); break
                                                                except: pass
                                                    else:
                                                        try: body = msg.get_payload(decode=True).decode('utf-8', errors='ignore')
                                                        except: pass
                                                    
                                                    clipped_body = body[:1000] + ("\n\n...(內容過長已省略，請至 Gmail 查看全文)" if len(body)>1000 else "")
                                                        
                                                    projects_db[proj_name]["replies"].append({
                                                        "company": company_name,
                                                        "email": addr,
                                                        "subject": subject_check,
                                                        "body": clipped_body,
                                                        "read": False,
                                                        "time": time.strftime("%Y-%m-%d %H:%M")
                                                    })
                                                    new_reply_count += 1
                                                    
                                                    mail.copy(num, reply_folder)
                                                    mail.store(num, '+FLAGS', '\\Deleted')
                        
                        mail.expunge()
                        
                        if new_reply_count > 0:
                            save_projects(projects_db)
                            st.success(f"🎉 成功攔截 {new_reply_count} 封新回信！")
                            time.sleep(2)
                            st.rerun()
                        else:
                            st.info("掃描完成，但沒有發現『未處理過』的目標廠商回信。")
                            
                        with st.expander("🛠️ 工程師除錯資訊 (點我展開)"):
                            st.warning(f"**資料庫裡記錄要找的目標 Email 有 {len(target_emails)} 個：**")
                            st.write(target_emails)
                            st.warning(f"**系統剛剛在信箱裡掃描到的最近幾封信來源是：**")
                            st.write(scanned_emails[-10:])
                    else:
                        st.info("近 3 天內沒有收到任何信件。")
                    mail.logout()
                except Exception as e:
                    st.error(f"連線或讀取失敗：{e}")

    st.divider()
    st.subheader("2. 廠商回信匣")
    has_any_reply = False
    
    for p_name, p_data in projects_db.items():
        replies = p_data.get("replies", [])
        if not replies: continue
        
        has_any_reply = True
        st.markdown(f"#### 📂 專案：{p_name}")
        
        for idx, reply in enumerate(reversed(replies)):
            real_idx = len(replies) - 1 - idx
            is_unread = not reply.get("read", True)
            icon = "🔴" if is_unread else "🟢"
            
            with st.expander(f"{icon} 來自 {reply['company']} ({reply['email']}) - {reply['subject']}"):
                st.caption(f"接收時間：{reply.get('time', '未知')}")
                st.text(reply['body'])
                
                if is_unread:
                    if st.button("標示為已讀", key=f"read_{p_name}_{real_idx}"):
                        projects_db[p_name]["replies"][real_idx]["read"] = True
                        save_projects(projects_db)
                        st.rerun()
    
    if not has_any_reply:
        st.info("目前資料庫中尚無任何廠商回信紀錄。")


# ==========================================
# 模式 C：專案與寄信區 (整合外網異常擷取偵測)
# ==========================================
elif "🏠 專案與寄信區" in app_mode:
    if st.session_state.page == "home":
        st.title("✉️ 公關寄信系統 - 專案大廳")
        col_new, col_list = st.columns([1, 2])
        
        with col_new:
            st.subheader("➕ 建立新專案")
            with st.form("new_project_form"):
                new_proj_name = st.text_input("專案命名：").strip()
                if st.form_submit_button("建立專案", type="primary", use_container_width=True):
                    if not new_proj_name: st.error("名稱不能為空！")
                    elif new_proj_name in projects_db: st.error("專案已存在！")
                    else:
                        projects_db[new_proj_name] = {"sent_companies": [], "template": DEFAULT_TEMPLATE, "replies": []}
                        save_projects(projects_db)
                        log_security_event("PROJECT_CREATE", st.session_state.username, "SUCCESS", f"建立新專案: {new_proj_name}")
                        st.rerun()

        with col_list:
            st.subheader("📂 現有專案列表")
            for proj_name, proj_data in list(projects_db.items()):
                with st.container(border=True):
                    c1, c2, c3 = st.columns([6, 2, 2])
                    with c1:
                        st.markdown(f"#### {proj_name}")
                        st.caption(f"寄出 {len(proj_data.get('sent_companies', []))} 封")
                    with c2:
                        if st.button("📂 開啟", key=f"open_{proj_name}", use_container_width=True):
                            st.session_state.current_project = proj_name
                            st.session_state.page = "project"
                            st.rerun()
                    with c3:
                        if st.session_state.role == "admin" and st.button("🗑️ 刪除", key=f"del_{proj_name}", use_container_width=True):
                            del projects_db[proj_name]
                            save_projects(projects_db)
                            log_security_event("PROJECT_DELETE", st.session_state.username, "WARNING", f"刪除專案: {proj_name}")
                            st.rerun()

    elif st.session_state.page == "project":
        current_proj = st.session_state.current_project
        proj_data = projects_db[current_proj]

        if st.sidebar.button("🔙 返回專案大廳", type="primary", use_container_width=True):
            st.session_state.page = "home"
            st.rerun()
            
        st.sidebar.divider()
        st.sidebar.header("🔐 寄件帳號設定")
        sender_email = st.sidebar.text_input("團隊 Gmail 信箱", value=st.session_state.gmail_account).strip()
        sender_password = st.sidebar.text_input("應用程式密碼", value=st.session_state.gmail_password, type="password").strip()
        st.session_state.gmail_account, st.session_state.gmail_password = sender_email, sender_password
        
        st.sidebar.header("📂 系統設定")
        pdf_dir = st.sidebar.text_input("本機附件資料夾名稱", value="企劃書檔案").strip()
        backup_folder = st.sidebar.text_input("Gmail 寄件備份標籤", value="FRC_Sponsorship").strip()

        st.title(f"📁 專案：{current_proj}")
        st.divider()

        st.header("Step 1: 團隊與信件格式")
        col1, col2, col3 = st.columns(3)
        team_name = col1.text_input("團隊名稱", value="FRC 團隊名稱").strip()
        contact_person = col2.text_input("聯絡人", value=st.session_state.real_name).strip()
        contact_phone = col3.text_input("電話", value="0912-345-678").strip()

        email_template = st.text_area("✏️ 信件內容", value=proj_data.get("template", DEFAULT_TEMPLATE), height=200)
        if email_template != proj_data.get("template"):
            projects_db[current_proj]["template"] = email_template
            save_projects(projects_db)

        st.header("Step 2 & 3: 載入與寄出")
        uploaded_file = st.file_uploader("上傳名單 (.xlsx)", type=["xlsx"])
        if uploaded_file is not None:
            records = read_excel_data(uploaded_file)
            if records:
                company_list = list(set([r.get('企業／贊助單位', '') for r in records if r.get('企業／贊助單位')]))
                selected_company = st.selectbox("🔍 選擇廠商", company_list) if company_list else None
                
                if selected_company:
                    # 呼叫資安模組 security.py 檢查外網異常資料擷取頻率
                    is_suspicious, count, warn_msg = scraping_detector.record_access(
                        st.session_state.username or "Guest", threshold_per_minute=30
                    )
                    if is_suspicious:
                        st.warning(warn_msg)
                        log_security_event("SCRAPING_WARNING", st.session_state.username, "WARNING", warn_msg)

                    row_data = next((r for r in records if r.get('企業／贊助單位') == selected_company), {})
                    to_email = extract_email(row_data.get('聯絡資訊', ''))
                    
                    if to_email:
                        st.success(f"📧 **即將寄出至 (目標信箱)：** `{to_email}`")
                    else:
                        st.error("⚠️ **警告：** 在 Excel 中找不到此廠商的有效 Email，將無法寄送！")
                    
                    pdf_filename = f"{str(row_data.get('編號', '000')).zfill(3)}_{selected_company}_贊助企劃書.pdf"
                    pdf_path = os.path.join(pdf_dir, pdf_filename)
                    
                    format_dict = dict(row_data)
                    format_dict.update({"team_name": team_name, "contact_person": contact_person, "contact_phone": contact_phone, "pdf_filename": pdf_filename})
                    
                    preview_text = safe_format_template(email_template, format_dict)
                    st.text(preview_text)
                    
                    sent_list = proj_data.get("sent_companies", [])
                    is_sent = any(isinstance(r, dict) and r.get("company") == selected_company for r in sent_list) or (selected_company in sent_list)
                    
                    if st.button("🚀 再次寄送" if is_sent else "🚀 確定寄送", type="primary"):
                        if not sender_email or not sender_password or not to_email:
                            st.error("⚠️ 帳號未設定，或廠商無有效 Email！")
                        else:
                            with st.spinner('寄送信件中...'):
                                try:
                                    server = smtplib.SMTP("smtp.gmail.com", 587, timeout=10)
                                    server.starttls()
                                    server.login(sender_email, sender_password)
                                    
                                    msg = MIMEMultipart()
                                    msg['From'], msg['To'], msg['Subject'] = sender_email, to_email, f"【贊助合作邀請】{team_name} — 敬致 {selected_company}"
                                    msg.attach(MIMEText(preview_text, 'plain', 'utf-8'))
                                    
                                    if os.path.exists(pdf_path):
                                        with open(pdf_path, 'rb') as f:
                                            attach = MIMEApplication(f.read(), _subtype="pdf")
                                            attach.add_header('Content-Disposition', 'attachment', filename=pdf_filename)
                                            msg.attach(attach)
                                    
                                    server.send_message(msg)
                                    server.quit()
                                    
                                    try:
                                        imap = imaplib.IMAP4_SSL("imap.gmail.com", timeout=10)
                                        imap.login(sender_email, sender_password)
                                        if imap.select(backup_folder)[0] != 'OK': imap.create(backup_folder)
                                        imap.append(backup_folder, '\\Seen', imaplib.Time2Internaldate(time.time()), msg.as_bytes())
                                        imap.logout()
                                    except: pass
                                    
                                    if not is_sent:
                                        proj_data["sent_companies"].append({
                                            "company": selected_company,
                                            "sender": st.session_state.real_name,
                                            "email": to_email
                                        })
                                        save_projects(projects_db)
                                    
                                    log_security_event("MAIL_SEND", st.session_state.username, "SUCCESS", f"成功寄出企劃書至 {selected_company} ({to_email})")
                                    st.success("✅ 寄出成功！")
                                    time.sleep(1.5)
                                    st.rerun()
                                except Exception as e:
                                    log_security_event("MAIL_SEND_FAIL", st.session_state.username, "FAILURE", f"寄件失敗 ({selected_company}): {e}")
                                    st.error(f"❌ 寄件失敗：{e}")